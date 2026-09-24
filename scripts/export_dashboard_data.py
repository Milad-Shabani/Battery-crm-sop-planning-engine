#!/usr/bin/env python3
"""Export dashboard/data.js by reading directly from the Excel workbook
that `crm_sop_planning.cli run` produces (data/processed/SOP_CRM_Report.xlsx)
— the same business deliverable a CRM/ops person would open, not a
separate database. Tables are located by their styled header row (the
same fill color `load/excel_export.py` uses), so this reads the workbook's
actual structure rather than hardcoded cell ranges.

    python -m crm_sop_planning.cli run
    python scripts/export_dashboard_data.py
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

EXCEL_PATH = Path("data/processed/SOP_CRM_Report.xlsx")
OUT_FILE = Path("dashboard/data.js")
HEADER_FILL_RGB = "001F2A44"  # openpyxl stores ARGB; matches HEADER_FILL in excel_export.py


def _extract_tables(ws) -> list[list[dict]]:
    """Find every header-styled row in a sheet and read the table beneath
    it until a blank row. Returns one list-of-dicts per table found."""
    tables = []
    row = 1
    max_row = ws.max_row
    while row <= max_row:
        cell = ws.cell(row=row, column=1)
        is_header = cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb == HEADER_FILL_RGB
        if is_header:
            headers = []
            col = 1
            while ws.cell(row=row, column=col).value not in (None, ""):
                headers.append(str(ws.cell(row=row, column=col).value))
                col += 1

            records = []
            r = row + 1
            while r <= max_row and ws.cell(row=r, column=1).value not in (None, ""):
                record = {}
                for c, h in enumerate(headers, start=1):
                    val = ws.cell(row=r, column=c).value
                    record[h] = val.isoformat() if hasattr(val, "isoformat") else val
                records.append(record)
                r += 1
            tables.append(records)
            row = r
        else:
            row += 1
    return tables


def main() -> None:
    if not EXCEL_PATH.exists():
        raise SystemExit(
            f"{EXCEL_PATH} not found — run `python -m crm_sop_planning.cli run` first."
        )

    wb = load_workbook(EXCEL_PATH, data_only=True)

    exec_tables = _extract_tables(wb["Executive Summary"])
    monthly_summary = exec_tables[0]

    forecast_tables = _extract_tables(wb["Demand Forecast"])
    demand_detail, forecast_by_line = forecast_tables[0], forecast_tables[1]

    production_tables = _extract_tables(wb["Production Plan"])
    production_plan = production_tables[0]

    crm_tables = _extract_tables(wb["CRM Pipeline & Funnel"])
    lead_funnel, win_rate, lose_reasons = crm_tables[0], crm_tables[1], crm_tables[2]

    rep_tables = _extract_tables(wb["Sales Rep Leaderboard"])
    rep_leaderboard = rep_tables[0]

    at_risk_tables = _extract_tables(wb["At-Risk Accounts"])
    at_risk = [r for r in at_risk_tables[0] if r.get("At Risk")]

    mrp = _extract_tables(wb["Material Requirements"])[0]

    total_forecast = sum(r["Forecast Units"] for r in monthly_summary)
    total_allocated = sum(r["Allocated Units"] for r in monthly_summary)
    total_revenue_lost = sum(r["Revenue Lost Irr"] for r in monthly_summary)

    # Line scorecard: planned output against forecast, and how much of each
    # line's forecast comes from the open CRM pipeline rather than history.
    lines: dict[str, dict] = {}
    for r in demand_detail:
        line = lines.setdefault(
            r["Line Id"],
            {
                "line_id": r["Line Id"],
                "category": r["Category"],
                "forecast_units": 0,
                "pipeline_units": 0.0,
                "allocated_units": 0,
                "unmet_units": 0,
            },
        )
        line["pipeline_units"] += r["Pipeline Units"] or 0
    for r in production_plan:
        line = lines[r["Line Id"]]
        line["forecast_units"] += r["Forecast Units"]
        line["allocated_units"] += r["Allocated Units"]
        line["unmet_units"] += r["Unmet Units"]
    line_scores = sorted(lines.values(), key=lambda r: r["line_id"])
    for line in line_scores:
        f = line["forecast_units"]
        line["fill_rate"] = line["allocated_units"] / f if f else 1.0
        line["pipeline_share"] = line["pipeline_units"] / f if f else 0.0
        line["pipeline_units"] = round(line["pipeline_units"], 1)

    # SKUs the plan can't fully serve, largest shortfall first.
    names = {r["Product Id"]: (r["Product Name"], r["Category"]) for r in demand_detail}
    short_skus = sorted(
        (
            {
                "product_name": names.get(r["Product Id"], (r["Product Id"], ""))[0],
                "category": names.get(r["Product Id"], ("", r["Line Id"]))[1],
                "period_start": r["Period Start"],
                "forecast_units": r["Forecast Units"],
                "allocated_units": r["Allocated Units"],
                "unmet_units": r["Unmet Units"],
            }
            for r in production_plan
            if r["Unmet Units"]
        ),
        key=lambda r: r["unmet_units"],
        reverse=True,
    )

    leads_total = sum(r["Total"] for r in lead_funnel)
    leads_qualified = sum(r["Qualified"] for r in lead_funnel)
    opps_total = sum(r["Opportunities"] for r in win_rate)
    opps_won = sum(r["Won"] for r in win_rate)
    peak_materials = max(monthly_summary, key=lambda r: r["Materials Below Safety Stock"])

    payload = {
        "monthly_summary": monthly_summary,
        "forecast_by_line": forecast_by_line,
        "production_plan": production_plan,
        "lead_funnel": lead_funnel,
        "win_rate": win_rate,
        "lose_reasons": lose_reasons,
        "rep_leaderboard": sorted(rep_leaderboard, key=lambda r: r["Won Value Irr"], reverse=True)[
            :8
        ],
        "at_risk_count": len(at_risk),
        "at_risk_accounts": at_risk[:10],
        "lines": line_scores,
        "short_skus": short_skus[:8],
        "kpi": {
            "total_forecast": total_forecast,
            "total_allocated": total_allocated,
            "fill_rate": total_allocated / total_forecast if total_forecast else 0,
            "revenue_lost": total_revenue_lost,
            "revenue_planned": sum(r["Revenue Captured Irr"] for r in monthly_summary),
            "unmet_units": sum(r["Unmet Units"] for r in monthly_summary),
            "pipeline_units": round(sum(r["Pipeline Units"] or 0 for r in demand_detail), 1),
            "n_months": len(monthly_summary),
            "leads_total": leads_total,
            "lead_conversion": leads_qualified / leads_total if leads_total else 0,
            "opportunities": opps_total,
            "win_rate": opps_won / opps_total if opps_total else 0,
            "won_value": sum(r["Won Value Irr"] for r in win_rate),
            "peak_materials_below_ss": peak_materials["Materials Below Safety Stock"],
            "peak_materials_month": peak_materials["Period Start"],
            "materials_to_reorder": len(
                {r["Material Id"] for r in mrp if (r["Recommended Order Qty"] or 0) > 0}
            ),
        },
    }

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(
        "const DASHBOARD_DATA = " + json.dumps(payload, ensure_ascii=False, default=str) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT_FILE} from {EXCEL_PATH}")


if __name__ == "__main__":
    main()
